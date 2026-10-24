"""parser-B: openpyxl 3.1.5 replay of Evidence Role (column F) counts on the
'Evidence Lineage' sheet. Read-only. Usage: python3 -I replay_openpyxl.py OUT.json XLSX..."""
import sys, json, hashlib, warnings, zipfile, re, collections
import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.utils import get_column_letter

LIST11 = ["PRIMARY_DERIVATION","EXECUTABLE_RECEIPT","RAW_DATA","FORMAL_PROOF",
          "INDEPENDENT_RECONSTRUCTION","SUMMARY","NAVIGATION","BALLOT",
          "REGISTER_COPY","REVIEW_OF_REVIEW","GOVERNANCE_PRIMARY"]

def nonempty(v):
    # whitespace rule: None -> empty; otherwise str(v).strip() == "" -> empty (Python str.strip, Unicode whitespace)
    return v is not None and str(v).strip() != ""

def sha(b): return hashlib.sha256(b).hexdigest()

def locate_sheet(path, title):
    z = zipfile.ZipFile(path)
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    for s in re.findall(r"<sheet\b[^>]*/>", wb):
        nm = re.search(r'\bname="([^"]*)"', s).group(1)
        if nm == title:
            rid = re.search(r'\br:id="([^"]*)"', s).group(1)
            for r in re.findall(r"<Relationship\b[^>]*/>", rels):
                if re.search(r'\bId="%s"' % re.escape(rid), r):
                    tgt = re.search(r'\bTarget="([^"]*)"', r).group(1)
                    member = tgt.lstrip("/") if tgt.startswith("/") else "xl/" + tgt
                    return s, r, member, sha(z.read(member)), len(z.read(member))
    return None

out = []
for p in sys.argv[2:]:
    raw = open(p, "rb").read()
    rec = {"file": p, "bytes": len(raw), "sha256": sha(raw)}
    loc = locate_sheet(p, "Evidence Lineage")
    rec["sheet_xml"], rec["rel_xml"], rec["member"], rec["member_sha256"], rec["member_bytes"] = loc
    with warnings.catch_warnings(record=True) as wl:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(p, read_only=False, data_only=False)
    rec["openpyxl_warnings"] = sorted(set(f"{w.category.__name__}: {w.message}" for w in wl))
    rec["sheetnames_index"] = wb.sheetnames.index("Evidence Lineage")
    ws = wb["Evidence Lineage"]
    rec["max_row"], rec["max_column"] = ws.max_row, ws.max_column
    rec["header_row"] = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rec["header_F"] = ws["F1"].value
    # validations
    dvs = []
    for dv in ws.data_validations.dataValidation:
        mcr = MultiCellRange(str(dv.sqref))
        touchesF = any(cr.min_col <= 6 <= cr.max_col for cr in mcr.ranges)
        f1 = dv.formula1
        kind = None; values = None
        if f1 is not None:
            if f1.startswith('"') and f1.endswith('"'):
                kind = "inline_list"; values = f1[1:-1].split(",")
            else:
                kind = "reference_or_formula"
        dvs.append({"sqref": str(dv.sqref), "touches_F": touchesF, "type": dv.type,
                    "operator": dv.operator, "formula1": f1, "formula2": dv.formula2,
                    "formula1_kind": kind, "inline_values": values,
                    "allowBlank": dv.allow_blank, "showErrorMessage": dv.showErrorMessage,
                    "showInputMessage": dv.showInputMessage, "showDropDown(OOXML:true=hide)": dv.showDropDown,
                    "errorStyle": dv.errorStyle, "promptTitle": dv.promptTitle, "prompt": dv.prompt,
                    "errorTitle": dv.errorTitle, "error": dv.error, "imeMode": dv.imeMode})
    rec["validations"] = dvs
    fdvs = [d for d in dvs if d["touches_F"]]
    covered = set()
    for d in fdvs:
        for cr in MultiCellRange(d["sqref"]).ranges:
            if cr.min_col <= 6 <= cr.max_col:
                covered.update(range(cr.min_row, cr.max_row + 1))
    # data scan
    rows_any = []; populated = []; blankF_other = []; formulas_F = []; types_F = collections.Counter()
    last_row_any = 0
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        if any(nonempty(v) for v in vals):
            rows_any.append(r); last_row_any = r
        f = vals[5]
        if f is not None:
            types_F[type(f).__name__] += 1
            if isinstance(f, str) and f.startswith("="): formulas_F.append((r, f))
        if nonempty(f):
            populated.append((r, f))
        elif any(nonempty(v) for v in vals[:5]):
            blankF_other.append(r)
    # last row with any content anywhere in the sheet (all columns)
    last_any_all = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            if nonempty(c.value): last_any_all = max(last_any_all, c.row)
    lab_raw = collections.Counter(str(v) for _, v in populated)
    lab_strip = collections.Counter(str(v).strip() for _, v in populated)
    in_raw = sum(n for k, n in lab_raw.items() if k in LIST11)
    in_strip = sum(n for k, n in lab_strip.items() if k in LIST11)
    in_ci = sum(n for k, n in lab_strip.items() if k.upper() in LIST11)
    dv_list = fdvs[0]["inline_values"] if fdvs and fdvs[0]["inline_values"] else None
    outside = collections.Counter({k: n for k, n in lab_strip.items() if k not in LIST11})
    rec.update({
        "rows_with_any_A_to_F": len(rows_any), "first_row_any": rows_any[0] if rows_any else None,
        "last_row_any_A_to_F": last_row_any, "last_row_any_any_column": last_any_all,
        "populated_F": len(populated),
        "in_list_raw_exact": in_raw, "in_list_after_strip": in_strip, "in_list_case_insensitive_after_strip": in_ci,
        "outside_list_after_strip": len(populated) - in_strip,
        "outside_list_raw_exact": len(populated) - in_raw,
        "distinct_nonempty_labels_raw": len(lab_raw), "distinct_nonempty_labels_strip": len(lab_strip),
        "labels_with_lead_trail_ws": sorted({k for k in lab_raw if k != k.strip()}),
        "blankF_rows_with_other_A_to_E": len(blankF_other), "blankF_rows_list": blankF_other,
        "F_value_types": dict(types_F), "F_formulas": formulas_F[:20], "n_F_formulas": len(formulas_F),
        "dv_list_equals_LIST11": dv_list == LIST11,
        "in_list_by_label": {k: lab_strip.get(k, 0) for k in LIST11},
        "top10_outside": outside.most_common(10),
        "all_outside_labels": outside.most_common(),
        "F_rows_populated_not_covered": [r for r, _ in populated if r not in covered],
        "data_rows_not_covered": [r for r in rows_any if r not in covered],
        "validation_covered_rows_min_max": [min(covered), max(covered)] if covered else None,
    })
    # first & last populated rows for citing
    rec["first_populated_F"] = populated[0] if populated else None
    rec["last_populated_F"] = populated[-1] if populated else None
    out.append(rec)

json.dump(out, open(sys.argv[1], "w"), indent=1, default=str)
for rec in out:
    print("====", rec["file"].split("/")[-1], rec["sha256"])
    for k in ["member","member_sha256","member_bytes","sheetnames_index","openpyxl_warnings","max_row","max_column","header_F",
              "rows_with_any_A_to_F","first_row_any","last_row_any_A_to_F","last_row_any_any_column","populated_F",
              "in_list_raw_exact","in_list_after_strip","in_list_case_insensitive_after_strip",
              "outside_list_after_strip","outside_list_raw_exact","distinct_nonempty_labels_raw","distinct_nonempty_labels_strip",
              "labels_with_lead_trail_ws","blankF_rows_with_other_A_to_E","blankF_rows_list","F_value_types","n_F_formulas",
              "dv_list_equals_LIST11","in_list_by_label","top10_outside","F_rows_populated_not_covered","data_rows_not_covered",
              "validation_covered_rows_min_max","first_populated_F","last_populated_F"]:
        print(f"  {k}: {rec[k]}")
    print("  header_row:", rec["header_row"])
    for d in rec["validations"]:
        print("  DV:", d)
